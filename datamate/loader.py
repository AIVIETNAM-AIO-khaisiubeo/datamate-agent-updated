"""Ingest one business package (CSV / XLSX / SQLite + README.md) into DuckDB.

Handles the messiness documented in the challenge data (rulebook §8 + READMEs):
  * report-style title rows above the header, two-row headers whose upper row
    holds group labels ("Tiền", "Đơn", "Hàng hoá"...);
  * one sheet per month and/or one file per month / per year -> sheets and files
    with the same column set are concatenated into ONE table (with `_sheet`);
  * values: VND money ('45k', '45.000đ', '6,719.0000'), quantities with units
    ('1 ly'), dates in mixed formats inside one column ('01/02/2025', Excel
    serial 45689, '1 thg 6 2025', '1/6' with the year taken from the sheet,
    ISO datetimes), booleans (TRUE/1, Có/Không);
  * 'Tổng cộng' total rows, repeated header rows, fully empty rows.
Group-level forward fill and duplicate-line removal need column roles and are
done in salesview.py.

Execution engine: DuckDB (Raasveldt & Mühleisen, SIGMOD 2019). Parsed tables
are cached on disk (pickle) keyed by file size/mtime + loader version.
"""
from __future__ import annotations

import datetime as dt
import glob
import hashlib
import os
import pickle
import re
import sqlite3
import warnings
from dataclasses import dataclass, field

import duckdb
import numpy as np
import pandas as pd

from .textnorm import fold, nfc, parse_number_vn, parse_qty_vn

warnings.filterwarnings("ignore")
LOADER_VERSION = "13"


def is_text(s: pd.Series) -> bool:
    return (pd.api.types.is_object_dtype(s) or pd.api.types.is_string_dtype(s)) and \
        not pd.api.types.is_datetime64_any_dtype(s)


def ident(name: str, fallback: str = "col") -> str:
    k = re.sub(r"[^0-9a-z]+", "_", fold(name)).strip("_")
    if not k:
        k = fallback
    if k[0].isdigit():
        k = f"c_{k}"
    return k


@dataclass
class Table:
    name: str
    df: pd.DataFrame
    source: str
    orig: dict = field(default_factory=dict)      # clean col -> original header
    notes: list = field(default_factory=list)
    files: list = field(default_factory=list)     # source file basenames (for README lookup)
    title: str = ""                               # original (un-normalised) sheet / file / SQL table name
    is_sql: bool = False


# ------------------------------------------------------------------ cells ---
def _s(v) -> str:
    if v is None or v is pd.NaT:
        return ""
    if isinstance(v, float) and np.isnan(v):
        return ""
    if not isinstance(v, str):
        try:
            if pd.isna(v):
                return ""
        except (TypeError, ValueError):
            pass
    return str(v).strip()


_YEAR_RE = re.compile(r"(?<!\d)(20\d{2})(?!\d)")
_MONTH_IN_NAME = [re.compile(p) for p in (r"(?:^|[^a-z0-9])t(?:hang)?[ _.-]*(\d{1,2})(?:[^0-9]|$)",
                                         r"(?<!\d)(\d{1,2})[ ._/-](?:20\d{2})", r"(?:20\d{2})[ ._/-](\d{1,2})(?!\d)")]


def year_month_hint(*names) -> tuple[int | None, int | None]:
    y = m = None
    for n in names:
        f = fold(n).replace(" ", "_") if n else ""
        if y is None:
            mm = _YEAR_RE.search(f)
            if mm:
                y = int(mm.group(1))
        if m is None:
            for rx in _MONTH_IN_NAME:
                mm = rx.search(f)
                if mm and 1 <= int(mm.group(1)) <= 12:
                    m = int(mm.group(1))
                    break
    return y, m


_D_DMY = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?)?$")
_D_YMD = re.compile(r"^(\d{4})[/.-](\d{1,2})[/.-](\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?(?:\.\d+)?)?$")
_D_THG = re.compile(r"^(\d{1,2})\s*(?:thg|th|thang|tháng)\s*(\d{1,2})[ ,/-]*(\d{4})?$")
_D_DM = re.compile(r"^(\d{1,2})[/.-](\d{1,2})$")


def parse_date_cell(v, year_hint=None, allow_serial=True):
    """-> pd.Timestamp or None."""
    if v is None:
        return None
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v
    if isinstance(v, (dt.datetime, dt.date)):
        return pd.Timestamp(v)
    if isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool):
        if np.isnan(v):
            return None
        if allow_serial and 20000 <= float(v) <= 80000:
            return pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(v))
        return None
    s = nfc(str(v)).strip().lower()
    if not s:
        return None
    if s.isdigit() and allow_serial and 20000 <= int(s) <= 80000:
        return pd.Timestamp("1899-12-30") + pd.Timedelta(days=int(s))
    try:
        m = _D_YMD.match(s)
        if m:
            y, mo, d, hh, mi, ss = m.groups()
            return pd.Timestamp(int(y), int(mo), int(d), int(hh or 0), int(mi or 0), int(ss or 0))
        m = _D_DMY.match(s)
        if m:
            d, mo, y, hh, mi, ss = m.groups()
            y = int(y) + (2000 if len(y) == 2 else 0)
            return pd.Timestamp(y, int(mo), int(d), int(hh or 0), int(mi or 0), int(ss or 0))
        m = _D_THG.match(s)
        if m:
            d, mo, y = m.groups()
            y = int(y) if y else year_hint
            return pd.Timestamp(int(y), int(mo), int(d)) if y else None
        m = _D_DM.match(s)
        if m and year_hint:
            return pd.Timestamp(int(year_hint), int(m.group(2)), int(m.group(1)))
    except (ValueError, TypeError):
        return None
    return None


_TIME_RE = re.compile(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$")


def parse_time_cell(v):
    if isinstance(v, dt.time):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, (dt.datetime, pd.Timestamp)):
        return v.hour * 3600 + v.minute * 60 + v.second
    m = _TIME_RE.match(_s(v))
    if m:
        return int(m.group(1)) * 3600 + int(m.group(2)) * 60 + int(m.group(3) or 0)
    return None


DATE_NAME = re.compile(r"\b(ngay|date|thoi gian|thoi diem|time|han|hsd|exp|expiry|nsx|tao|created|lap|nhan|tra|giao|sinh|dang ky|ap dung)\b")
QTY_NAME = re.compile(r"\b(sl|s l|so luong|qty|quantity|slg|so|soluong|sl quy doi|quy ra|so ly)\b")
CODE_LIKE = re.compile(r"^[A-Za-z]{1,6}[-_]?\d{2,}$|^0\d{6,}$")


# ------------------------------------------------------------ header ---
def _header_score(row) -> float:
    vals = [_s(v) for v in row]
    ne = [v for v in vals if v]
    if not ne:
        return -1.0
    texty = [v for v in ne if parse_number_vn(v) is None and parse_date_cell(v, 2025, allow_serial=False) is None
             and not isinstance(v, (dt.datetime, dt.date))]
    return len(set(texty)) + 0.01 * len(ne)


def fix_header(raw: pd.DataFrame, notes: list) -> pd.DataFrame:
    raw = raw.dropna(how="all").dropna(axis=1, how="all")
    if raw.empty:
        return raw
    raw = raw.reset_index(drop=True)
    n = min(15, len(raw))
    scores = [_header_score(raw.iloc[i]) for i in range(n)]
    best = max(scores)
    h = next(i for i in range(n) if scores[i] >= best - 1e-9)
    header = [_s(v) for v in raw.iloc[h]]
    # group-label row directly above (e.g. 'Tiền' over 'ĐG','TT'): only used to fill blank header cells
    if h > 0 and any(not c for c in header):
        above = [_s(v) for v in raw.iloc[h - 1]]
        last = ""
        for i, (a, c) in enumerate(zip(above, header)):
            last = a or last
            if not c:
                header[i] = last or f"col{i + 1}"
    if h > 0:
        notes.append(f"header at row {h + 1}")
    df = raw.iloc[h + 1:].copy()
    cols, seen = [], {}
    for i, c in enumerate(header):
        c = nfc(c) or f"col{i + 1}"
        seen[c] = seen.get(c, 0) + 1
        cols.append(c if seen[c] == 1 else f"{c} {seen[c]}")
    df.columns = cols
    if df.empty:
        return df
    # footnotes ('Ghi chú: tiền tính bằng đồng') : rows with a single non-empty cell in a wide table
    if len(cols) >= 4:
        nonempty = sum((df[c].notna() & (df[c].astype(str).str.strip() != "")).astype(int) for c in df.columns)
        df = df[nonempty >= 2]
    # drop repeated header rows (concatenated exports) and fully empty rows
    if len(df):
        hits = sum((df[c].astype(str).str.strip().str.lower() == str(c).strip().lower()).astype(int)
                   for c in df.columns)
        rep = hits >= max(2, 0.6 * len(cols))
        if rep.any():
            df = df[~rep]
    return df.dropna(how="all")


TOTAL_WORDS = {"tong", "tong cong", "cong", "total", "grand total", "tong so", "cong don", "tong ket"}


def drop_total_rows(df: pd.DataFrame, notes: list) -> pd.DataFrame:
    """Drop subtotal / total rows. The label must be the row's FIRST non-empty cell, so a
    customer called "Tống Nam Minh" in a later column never looks like a "tổng ..." row."""
    if df.empty:
        return df
    cols = [c for c in df.columns if c != "_sheet"][:20]
    if not cols:
        return df
    first = pd.Series("", index=df.index)
    for c in reversed(cols):     # first non-empty TEXT cell (numbers such as the subtotal amount are skipped)
        if pd.api.types.is_numeric_dtype(df[c]) or pd.api.types.is_datetime64_any_dtype(df[c]):
            continue
        v = df[c].map(_s)
        v = v.where(~v.str.match(r"^[-+]?[\d.,\s]+(đ|d|vnd)?$"), "")
        first = v.where(v != "", first)
    f = first.map(fold)
    sub = r"^(tong|cong|tong cong|subtotal|total|grand total)( (ngay|thang|tuan|ky|quy|nam|ca|so|tien|cong|doanh thu|phat sinh)\b.*|\s*:.*|\s+\d.*)?$"
    mask = f.isin(TOTAL_WORDS) | f.str.match(sub)
    if mask.any() and mask.sum() <= max(3, 0.05 * len(df)):
        notes.append(f"dropped {int(mask.sum())} total row(s)")
        return df[~mask]
    return df


# ------------------------------------------------------------- coerce ---
def name_key(name: str) -> str:
    """'NgayLap' / 'ngay_lap' / 'Ngày lập' -> 'ngay lap' (folded, camelCase split)."""
    src = nfc(str(name))
    out = []
    for i, ch in enumerate(src):
        if i and ch.isupper() and (src[i - 1].islower() or
                                   (src[i - 1].isupper() and i + 1 < len(src) and src[i + 1].islower())):
            out.append(" ")
        out.append(ch)
    return fold("".join(out).replace("_", " ").replace(".", " "))


def coerce_column(name: str, s: pd.Series, year_hints: pd.Series | None, notes: list) -> pd.Series:
    nm = name_key(name)
    if pd.api.types.is_datetime64_any_dtype(s) or pd.api.types.is_bool_dtype(s):
        return s
    nonnull = s[s.map(lambda v: _s(v) != "")]
    if nonnull.empty:
        return s.map(lambda v: None)
    date_hint = bool(DATE_NAME.search(nm))
    sample = nonnull if len(nonnull) <= 3000 else nonnull.sample(3000, random_state=0)
    yh_sample = year_hints.loc[sample.index] if year_hints is not None else None

    # --- dates (mixed formats, Excel serials only if the column name suggests a date)
    def pd_(v, i):
        return parse_date_cell(v, None if yh_sample is None else yh_sample.get(i), allow_serial=date_hint)
    parsed = [pd_(v, i) for i, v in sample.items()]
    ok = sum(p is not None for p in parsed) / len(sample)
    numeric_raw = sample.map(lambda v: isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool)).mean()
    if ok > 0.9 and (date_hint or numeric_raw < 0.5):
        yh = year_hints if year_hints is not None else pd.Series(None, index=s.index)
        out = pd.Series([parse_date_cell(v, yh.get(i), allow_serial=date_hint) for i, v in s.items()], index=s.index)
        out = pd.to_datetime(out, errors="coerce")
        if date_hint is False:
            notes.append(f"{name}: dates")
        return out
    # --- pure time-of-day column ('Giờ')
    if re.search(r"\b(gio|time)\b", nm):
        tt = sample.map(parse_time_cell)
        if tt.notna().mean() > 0.9:
            return s.map(lambda v: (lambda x: None if x is None else f"{x // 3600:02d}:{x % 3600 // 60:02d}:{x % 60:02d}")(parse_time_cell(v)))
    # --- booleans
    strs = sample.map(_s)
    fb = strs.map(fold)
    bset_t, bset_f = {"true", "co", "yes", "x", "1", "dung"}, {"false", "khong", "no", "0", "sai"}
    words = (bset_t | bset_f) - {"0", "1"}
    if fb.isin(bset_t | bset_f).mean() > 0.97 and fb.isin(words).any():
        return s.map(lambda v: None if _s(v) == "" else fold(_s(v)) in bset_t)
    # --- numbers / money / quantities
    looks_code = strs.str.match(CODE_LIKE.pattern).mean() > 0.3
    if not looks_code:
        qty_like = bool(QTY_NAME.search(nm))
        parser = parse_qty_vn if qty_like else parse_number_vn
        nums = sample.map(parser)
        if nums.notna().mean() > 0.97 or (qty_like and nums.notna().mean() > 0.9):
            out = s.map(lambda v: parser(v) if _s(v) != "" else None)
            out = pd.to_numeric(out, errors="coerce")
            return out
        if fb.isin(bset_t | bset_f).mean() > 0.97 and fb.nunique() <= 2:
            return s.map(lambda v: None if _s(v) == "" else fold(_s(v)) in bset_t)
    return s.map(lambda v: nfc(_s(v)) if _s(v) != "" else None)


# ------------------------------------------------------------ readers ---
def _read_csv(p):
    for enc in ("utf-8-sig", "utf-8", "cp1258", "latin1"):
        try:
            return pd.read_csv(p, header=None, dtype=str, encoding=enc, sep=None, engine="python",
                               keep_default_na=False, na_values=[""], skip_blank_lines=True)
        except UnicodeDecodeError:
            continue
        except Exception:
            try:
                return pd.read_csv(p, header=None, dtype=str, encoding=enc, keep_default_na=False,
                                   na_values=[""], on_bad_lines="skip")
            except Exception:
                continue
    return None


def _read_xlsx(p):
    try:  # Rust reader (python-calamine): ~10x faster than openpyxl on large exports
        sheets = pd.read_excel(p, sheet_name=None, header=None, dtype=object, engine="calamine")
        return [(name, df) for name, df in sheets.items() if not df.empty]
    except Exception:
        pass
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            out.append((ws.title, pd.DataFrame(rows, dtype=object)))
    wb.close()
    return out


def read_package(biz_dir: str):
    """-> list of raw parts: (logical_name, file_basename, sheet, DataFrame(header resolved, raw values), notes, is_sql)."""
    parts = []
    for p in sorted(glob.glob(os.path.join(biz_dir, "**", "*"), recursive=True)):
        base, ext = os.path.splitext(os.path.basename(p))
        ext = ext.lower()
        fname = os.path.basename(p)
        if base.lower() == "readme" or os.path.isdir(p):
            continue
        if ext in (".csv", ".tsv", ".txt"):
            raw = _read_csv(p)
            if raw is not None:
                notes = []
                parts.append((base, fname, "", fix_header(raw, notes), notes, False))
        elif ext in (".xlsx", ".xlsm", ".xls"):
            try:
                sheets = _read_xlsx(p)
            except Exception:
                sheets = list(pd.read_excel(p, sheet_name=None, header=None, dtype=object).items())
            for sh, raw in sheets:
                notes = []
                df = fix_header(raw, notes)
                if not df.empty and len(df.columns) >= 1:
                    parts.append((base, fname, sh, df, notes, False))
        elif ext in (".sqlite", ".sqlite3", ".db"):
            c = sqlite3.connect(f"file:{p}?mode=ro", uri=True)
            for (t,) in c.execute("SELECT name FROM sqlite_master WHERE type IN ('table','view') AND name NOT LIKE 'sqlite_%'"):
                df = pd.read_sql_query(f'SELECT * FROM "{t}"', c)
                parts.append((t, f"{fname}:{t}", "", df, ["from SQLite"], True))
            c.close()
    return parts


def _logical_name(base: str) -> str:
    """'SoBanHang_2026' / 'BanHang_T1_2025' / 'DanhSachHoaDonChiTiet_2' -> shared stem."""
    s = re.sub(r"([_ -]?(t|thang)?\d{1,2})?[_ -]?20\d{2}$", "", base, flags=re.I)
    s = re.sub(r"[_ -]\d{1,2}$", "", s)
    return s or base


# ----------------------------------------------------------- business ---
class Business:
    def __init__(self, biz_dir: str, business_id: str | None = None, cache_dir: str | None = None):
        self.dir = biz_dir
        self.id = business_id or os.path.basename(os.path.normpath(biz_dir))
        self.con = duckdb.connect(":memory:")
        self.tables: dict[str, Table] = {}
        self.readme = ""
        self.load_errors: list[str] = []
        rd = [p for p in glob.glob(os.path.join(biz_dir, "**", "*"), recursive=True)
              if os.path.basename(p).lower() in ("readme.md", "readme.txt", "readme")]
        if rd:
            self.readme = open(rd[0], encoding="utf-8", errors="replace").read()
        cached = self._cache_load(cache_dir)
        if cached is None:
            self._load()
            self._cache_save(cache_dir)
        for t in self.tables.values():
            self._to_duckdb(t)

    # ---- cache --------------------------------------------------------
    def _cache_key(self):
        h = hashlib.sha1(LOADER_VERSION.encode())
        for p in sorted(glob.glob(os.path.join(self.dir, "**", "*"), recursive=True)):
            if os.path.isfile(p):
                st = os.stat(p)
                h.update(f"{os.path.relpath(p, self.dir)}:{st.st_size}:{int(st.st_mtime)}".encode())
        return h.hexdigest()[:16]

    def _cache_load(self, cache_dir):
        if not cache_dir:
            return None
        p = os.path.join(cache_dir, "tables", f"{self.id}_{self._cache_key()}.pkl")
        if os.path.exists(p):
            try:
                with open(p, "rb") as f:
                    d = pickle.load(f)
                self.tables, self.load_errors = d["tables"], d["errors"]
                return True
            except Exception:
                return None
        return None

    def _cache_save(self, cache_dir):
        if not cache_dir:
            return
        try:
            os.makedirs(os.path.join(cache_dir, "tables"), exist_ok=True)
            p = os.path.join(cache_dir, "tables", f"{self.id}_{self._cache_key()}.pkl")
            with open(p, "wb") as f:
                pickle.dump({"tables": self.tables, "errors": self.load_errors}, f)
        except Exception:
            pass

    # ---- loading -------------------------------------------------------
    def _unique(self, name: str) -> str:
        base, i = name, 2
        while name in self.tables:
            name, i = f"{base}_{i}", i + 1
        return name

    def _load(self):
        try:
            parts = read_package(self.dir)
        except Exception as e:
            self.load_errors.append(f"read: {type(e).__name__}: {e}")
            return
        sheets_per_file: dict = {}
        for _, fname, _, _, _, _ in parts:
            sheets_per_file[fname] = sheets_per_file.get(fname, 0) + 1
        # group parts with identical column sets (monthly sheets / yearly files)
        groups: dict = {}
        for base, fname, sh, df, notes, is_sql in parts:
            key = (is_sql, tuple(sorted(ident(str(c)) for c in df.columns)), base if is_sql else "")
            groups.setdefault(key, []).append((base, fname, sh, df, notes))
        for (is_sql, _, _), items in groups.items():
            try:
                generic = lambda sh: (not sh) or re.fullmatch(r"sheet\d*", fold(sh).replace(" ", "")) is not None
                if len(items) > 1:
                    bases = {it[0] for it in items}
                    sheets = {it[2] for it in items}
                    if len(sheets) == 1 and not generic(next(iter(sheets))) and len(bases) > 1:
                        name = next(iter(sheets))        # same sheet in several yearly workbooks
                    elif len(bases) == 1:
                        name = items[0][0]               # monthly sheets of one workbook
                    else:
                        name = _logical_name(items[0][0])
                else:
                    base, fname, sh = items[0][:3]
                    name = base if (is_sql or generic(sh) or sheets_per_file.get(fname, 1) == 1) else sh
                self._register_group(items, is_sql, name)
            except Exception as e:
                self.load_errors.append(f"{items[0][1]}: {type(e).__name__}: {e}"[:300])

    def _register_group(self, items, is_sql, base_name):
        frames, files, notes = [], [], []
        multi = len(items) > 1
        refmap = {ident(str(c)): c for c in items[0][3].columns}
        for base, fname, sh, df, nt in items:
            df = df.copy()
            df.columns = [refmap.get(ident(str(c)), c) for c in df.columns]
            y, _ = year_month_hint(sh, base)
            df["__year"] = y
            if multi:
                df["_sheet"] = sh or base
            frames.append(df)
            files.append(fname)
            notes += nt
        df = pd.concat(frames, ignore_index=True) if multi else frames[0].reset_index(drop=True)
        name = self._unique(ident(base_name, "t"))
        year_hints = df.pop("__year")
        orig, cols, seen = {}, [], {}
        for c in df.columns:
            k = ident(str(c))
            seen[k] = seen.get(k, 0) + 1
            k = k if seen[k] == 1 else f"{k}_{seen[k]}"
            orig[k] = nfc(str(c))
            cols.append(k)
        df.columns = cols
        if not is_sql:
            df = drop_total_rows(df, notes)
        year_hints = year_hints.loc[df.index] if year_hints is not None else None
        for c in list(df.columns):
            if c == "_sheet":
                continue
            df[c] = coerce_column(orig[c], df[c], year_hints, notes)
        df = df.reset_index(drop=True)
        if multi:
            notes.append(f"{len(items)} sheets/files concatenated")
        self.tables[name] = Table(name, df, ", ".join(sorted(set(files)))[:200] +
                                  (f" ({len(items)} parts)" if multi else ""), orig, sorted(set(notes)),
                                  sorted(set(f.split(":")[0] for f in files)), title=base_name, is_sql=is_sql)

    def _to_duckdb(self, t: Table):
        df = t.df
        self.con.register(f"__df_{t.name}", df)
        self.con.execute(f'CREATE OR REPLACE TABLE "{t.name}" AS SELECT * FROM "__df_{t.name}"')
        self.con.unregister(f"__df_{t.name}")

    def replace_table(self, name: str, df: pd.DataFrame):
        self.tables[name].df = df
        self._to_duckdb(self.tables[name])

    def add_view_table(self, name: str, df: pd.DataFrame, source: str, notes=None):
        self.tables[name] = Table(name, df, source, {c: c for c in df.columns}, notes or [], [])
        self._to_duckdb(self.tables[name])

    # ---- introspection -------------------------------------------------
    def date_range(self):
        """Activity date range (sales/orders), ignoring expiry / birth / creation dates."""
        lo = hi = None
        for t in self.tables.values():
            if t.name.startswith("v_"):
                continue
            for c in t.df.columns:
                s = t.df[c]
                if pd.api.types.is_datetime64_any_dtype(s) and s.notna().any() and not re.search(
                        r"han|hsd|exp|nsx|sinh|tao|dang_ky|vao_lam|khai_truong|ap_dung|thoi_viec|nhan_viec", c):
                    a, b = s.min(), s.max()
                    lo = a if lo is None or a < lo else lo
                    hi = b if hi is None or b > hi else hi
        return (lo.date() if lo is not None else None, hi.date() if hi is not None else None)

    def profile(self, max_vals: int = 10, roles: dict | None = None, skip_views: bool = False) -> str:
        lines = []
        for t in self.tables.values():
            if skip_views and t.name.startswith("v_"):
                continue
            df = t.df
            lines.append(f'TABLE "{t.name}"  -- {len(df)} rows, source: {t.source}'
                         + (f"  [{'; '.join(t.notes[:4])}]" if t.notes else ""))
            for c in df.columns:
                s = df[c]
                o = t.orig.get(c, c)
                role = (roles or {}).get((t.name, c))
                head = f"  - {c}" + (f' ("{o}")' if fold(o).replace(" ", "_") != c else "") + (f" <{role}>" if role else "")
                nn = int(s.notna().sum())
                try:
                    if pd.api.types.is_datetime64_any_dtype(s):
                        has_t = nn and (s.dropna().dt.hour.ne(0) | s.dropna().dt.minute.ne(0)).any()
                        info = f"{'DATETIME' if has_t else 'DATE'} [{s.min()} .. {s.max()}]" if nn else "DATE (empty)"
                    elif pd.api.types.is_bool_dtype(s) or (nn and s.dropna().map(type).eq(bool).all()):
                        info = f"BOOL true={int(s.fillna(False).astype(bool).sum())}"
                    elif pd.api.types.is_numeric_dtype(s):
                        info = (f"NUM [{s.min():g} .. {s.max():g}]" if nn else "NUM (empty)") + \
                               (" integer" if nn and (s.dropna() % 1 == 0).all() else "")
                    else:
                        v = s.dropna().astype(str)
                        u = v.unique()
                        ex = ", ".join(repr(x[:40]) for x in u[:max_vals])
                        info = f"TEXT {len(u)} distinct: {ex}{' ...' if len(u) > max_vals else ''}"
                except Exception:
                    info = "?"
                lines.append(f"{head}: {info}" + (f"  (nulls {len(df) - nn})" if nn < len(df) else ""))
        return "\n".join(lines)

    def snapshot(self, cache_dir: str) -> str:
        """Persist tables to a DuckDB file (read-only access for the Python sandbox)."""
        os.makedirs(cache_dir, exist_ok=True)
        path = os.path.abspath(os.path.join(cache_dir, f"{self.id}.duckdb"))
        if os.path.exists(path):
            os.remove(path)
        alias = "snap_" + re.sub(r"\W", "_", self.id)
        self.con.execute(f"ATTACH '{path}' AS {alias}")
        for name in self.tables:
            self.con.execute(f'CREATE TABLE {alias}."{name}" AS SELECT * FROM main."{name}"')
        self.con.execute(f"DETACH {alias}")
        return path
